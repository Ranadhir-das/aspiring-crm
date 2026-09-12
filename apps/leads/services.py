import csv
import io
from decimal import Decimal, InvalidOperation

from django.contrib.auth import get_user_model
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from .models import Lead, LeadAssignmentHistory, LeadImportBatch
from .utils import normalize_phone


REQUIRED_COLUMNS = {
    "name",
    "phone",
}


OPTIONAL_COLUMNS = {
    "email",
    "location",
    "college",
    "neet_status",
    "pcb_percentage",
    "preferred_intake",
    "source",
    "campaign",
    "notes",
}


def normalize_column_name(column):
    """
    Convert Excel/CSV column names into a consistent format.
    """

    if column is None:
        return ""

    return str(column).strip().lower().replace(" ", "_")


def read_csv_file(file):
    """
    Read a CSV file and return rows as dictionaries.
    """

    content = file.read()

    if isinstance(content, bytes):
        content = content.decode("utf-8-sig")

    reader = csv.DictReader(io.StringIO(content))

    headers = [
        normalize_column_name(header)
        for header in reader.fieldnames or []
    ]

    rows = []

    for row in reader:
        normalized_row = {}

        for key, value in row.items():
            normalized_key = normalize_column_name(key)

            normalized_row[normalized_key] = (
                value.strip()
                if isinstance(value, str)
                else value
            )

        rows.append(normalized_row)

    return headers, rows


def read_xlsx_file(file):
    """
    Read an XLSX file and return rows as dictionaries.
    """

    workbook = load_workbook(
        filename=file,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook.active

    rows_data = list(
        worksheet.iter_rows(values_only=True)
    )

    if not rows_data:
        return [], []

    headers = [
        normalize_column_name(header)
        for header in rows_data[0]
    ]

    rows = []

    for row in rows_data[1:]:
        normalized_row = {}

        for index, value in enumerate(row):
            if index >= len(headers):
                continue

            header = headers[index]

            if header:
                normalized_row[header] = (
                    str(value).strip()
                    if value is not None
                    else ""
                )

        rows.append(normalized_row)

    return headers, rows


def read_import_file(file):
    """
    Read CSV or XLSX file based on its extension.
    """

    filename = getattr(file, "name", "").lower()

    if filename.endswith(".csv"):
        return read_csv_file(file)

    if filename.endswith(".xlsx"):
        return read_xlsx_file(file)

    raise ValueError(
        "Unsupported file format. Please upload a CSV or XLSX file."
    )


def validate_columns(headers):
    """
    Validate required and unknown columns.
    """

    header_set = set(headers)

    missing_columns = REQUIRED_COLUMNS - header_set

    unknown_columns = (
        header_set
        - REQUIRED_COLUMNS
        - OPTIONAL_COLUMNS
    )

    return {
        "valid": not missing_columns,
        "missing_columns": sorted(missing_columns),
        "unknown_columns": sorted(unknown_columns),
    }


def validate_row(row, row_number):
    """
    Validate one imported lead row.
    """

    errors = []
    warnings = []

    name = str(row.get("name", "")).strip()
    phone = str(row.get("phone", "")).strip()

    if not name:
        errors.append("Name is required.")

    if not phone:
        errors.append("Phone is required.")

    normalized_phone = normalize_phone(phone)

    if phone and not normalized_phone:
        errors.append("Phone number is invalid.")

    if normalized_phone and len(normalized_phone) < 7:
        errors.append("Phone number is too short.")

    email = str(row.get("email", "")).strip()

    if email and "@" not in email:
        warnings.append(
            "Email format may be invalid."
        )

    # Validate PCB percentage if supplied.
    pcb_percentage = row.get("pcb_percentage")

    if pcb_percentage not in (None, ""):
        try:
            percentage = Decimal(
                str(pcb_percentage).strip()
            )

            if percentage < 0 or percentage > 100:
                errors.append(
                    "PCB percentage must be between 0 and 100."
                )

        except (InvalidOperation, ValueError):
            errors.append(
                "PCB percentage must be a valid number."
            )

    return {
        "row_number": row_number,
        "errors": errors,
        "warnings": warnings,
        "normalized_phone": normalized_phone,
    }


def get_existing_phone_map():
    """
    Load existing lead phone numbers once.

    Returns:
        dict[str, Lead]
    """

    phone_map = {}

    for lead in (
        Lead.objects
        .only("id", "name", "phone")
        .iterator()
    ):
        normalized = normalize_phone(lead.phone)

        if normalized:
            phone_map[normalized] = lead

    return phone_map


def check_duplicate_from_map(phone, phone_map):
    """
    Check for a duplicate using an in-memory phone map.
    """

    normalized = normalize_phone(phone)

    if not normalized:
        return None

    return phone_map.get(normalized)


def prepare_pcb_percentage(value):
    """
    Convert PCB percentage into a Decimal or None.

    Validation is performed separately by validate_row().
    """

    if value in (None, ""):
        return None

    return Decimal(str(value).strip())


def preview_import(file):
    """
    Read and validate an import file without creating leads.
    """

    headers, rows = read_import_file(file)

    column_result = validate_columns(headers)

    if not column_result["valid"]:
        return {
            "success": False,
            "columns": column_result,
            "total_rows": len(rows),
            "rows": [],
        }

    preview_rows = []

    valid_count = 0
    warning_count = 0
    duplicate_count = 0
    error_count = 0

    # Load existing database phones only once.
    existing_phone_map = get_existing_phone_map()

    # Track duplicate phones inside the uploaded file itself.
    seen_phones = set()

    for index, row in enumerate(rows, start=2):

        validation = validate_row(row, index)

        duplicate = None

        if not validation["errors"]:

            normalized_phone = validation["normalized_phone"]

            # Check duplicates already stored in database.
            duplicate = check_duplicate_from_map(
                normalized_phone,
                existing_phone_map,
            )

            # Check duplicates within the same uploaded file.
            if (
                not duplicate
                and normalized_phone in seen_phones
            ):
                validation["warnings"].append(
                    "Duplicate phone number in uploaded file."
                )

                duplicate = True

            if normalized_phone:
                seen_phones.add(normalized_phone)

        if duplicate:
            duplicate_count += 1

            if isinstance(duplicate, Lead):
                validation["warnings"].append(
                    f"Duplicate lead already exists: "
                    f"{duplicate.name} ({duplicate.phone})"
                )

        if validation["errors"]:
            error_count += 1
        else:
            valid_count += 1

        if validation["warnings"]:
            warning_count += 1

        preview_rows.append({
            "row_number": index,
            "data": row,
            "errors": validation["errors"],
            "warnings": validation["warnings"],
            "duplicate": bool(duplicate),
            "normalized_phone": validation["normalized_phone"],
        })

    return {
        "success": True,
        "columns": column_result,
        "total_rows": len(rows),
        "valid_count": valid_count,
        "warning_count": warning_count,
        "duplicate_count": duplicate_count,
        "error_count": error_count,
        "rows": preview_rows,
    }


def commit_import(file, imported_by):
    """
    Import valid, non-duplicate leads into the database.

    Returns a summary of the import.
    """

    headers, rows = read_import_file(file)

    column_result = validate_columns(headers)

    if not column_result["valid"]:
        raise ValueError(
            f"Missing required columns: "
            f"{', '.join(column_result['missing_columns'])}"
        )

    created_count = 0
    duplicate_count = 0
    warning_count = 0
    skipped_count = 0

    with transaction.atomic():

        batch = LeadImportBatch.objects.create(
            filename=file.name,
            imported_by=imported_by,
            total_rows=len(rows),
        )

        # Load existing database phones only once.
        existing_phone_map = get_existing_phone_map()

        # Track phones from the current upload.
        seen_phones = set()

        for index, row in enumerate(rows, start=2):

            validation = validate_row(row, index)

            if validation["errors"]:
                skipped_count += 1
                continue

            if validation["warnings"]:
                warning_count += 1

            normalized_phone = validation["normalized_phone"]

            # Check existing database leads.
            duplicate = check_duplicate_from_map(
                normalized_phone,
                existing_phone_map,
            )

            if duplicate:
                duplicate_count += 1
                continue

            # Check duplicate within the uploaded file.
            if normalized_phone in seen_phones:
                duplicate_count += 1
                continue

            seen_phones.add(normalized_phone)

            pcb_percentage = prepare_pcb_percentage(
                row.get("pcb_percentage")
            )

            lead = Lead.objects.create(
                name=str(
                    row.get("name", "")
                ).strip(),

                phone=str(
                    row.get("phone", "")
                ).strip(),

                email=str(
                    row.get("email", "")
                ).strip(),

                location=str(
                    row.get("location", "")
                ).strip(),

                college=str(
                    row.get("college", "")
                ).strip(),

                neet_status=str(
                    row.get("neet_status", "")
                ).strip(),

                pcb_percentage=pcb_percentage,

                preferred_intake=str(
                    row.get("preferred_intake", "")
                ).strip(),

                source=str(
                    row.get("source", "")
                ).strip(),

                campaign=str(
                    row.get("campaign", "")
                ).strip(),

                notes=str(
                    row.get("notes", "")
                ).strip(),
            )

            created_count += 1

            # Add the newly created lead to the map.
            # This prevents duplicate creation later in the
            # same import operation.
            existing_phone_map[normalized_phone] = lead

        batch.created_count = created_count
        batch.duplicate_count = duplicate_count
        batch.warning_count = warning_count

        batch.save(
            update_fields=[
                "created_count",
                "duplicate_count",
                "warning_count",
            ]
        )

    return {
        "success": True,
        "batch_id": batch.id,
        "filename": file.name,
        "total_rows": len(rows),
        "created_count": created_count,
        "duplicate_count": duplicate_count,
        "warning_count": warning_count,
        "skipped_count": skipped_count,
    }


def bulk_assign_leads(
    lead_ids,
    new_caller,
    assigned_by,
    reassign=False,
    reason="",
):
    """
    Assign multiple leads to a caller.

    If reassign=False:
        Already assigned leads are skipped.

    If reassign=True:
        Already assigned leads are reassigned.

    Returns a summary of the operation.
    """

    User = get_user_model()

    if new_caller.role != User.Role.CALLER:
        raise ValueError(
            "Selected user is not a caller."
        )

    if not lead_ids:
        raise ValueError(
            "No leads were selected."
        )

    assigned_count = 0
    skipped_count = 0
    reassigned_count = 0

    with transaction.atomic():

        leads = (
            Lead.objects
            .select_for_update()
            .filter(id__in=lead_ids)
        )

        for lead in leads:

            previous_caller = lead.assigned_caller

            # Already assigned and reassignment is disabled.
            if previous_caller and not reassign:
                skipped_count += 1
                continue

            # Already assigned to the same caller.
            if previous_caller == new_caller:
                skipped_count += 1
                continue

            lead.assigned_caller = new_caller
            lead.assigned_at = timezone.now()

            lead.save(
                update_fields=[
                    "assigned_caller",
                    "assigned_at",
                    "updated_at",
                ]
            )

            LeadAssignmentHistory.objects.create(
                lead=lead,
                previous_caller=previous_caller,
                new_caller=new_caller,
                assigned_by=assigned_by,
                reason=reason,
            )

            if previous_caller:
                reassigned_count += 1
            else:
                assigned_count += 1

    return {
        "success": True,
        "total_selected": len(lead_ids),
        "assigned_count": assigned_count,
        "reassigned_count": reassigned_count,
        "skipped_count": skipped_count,
    }