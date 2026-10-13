import csv
import io
from datetime import datetime
from django.http import HttpResponse, StreamingHttpResponse
from django.utils import timezone
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def sanitize_cell(value):
    """Protects against spreadsheet formula injection (CSV/Excel) from user-controlled strings.
    If a string starts with formula trigger characters (=, +, -, @, \t, \r), prepends a single quote.
    """
    if value is None:
        return ""
    if isinstance(value, str):
        if not value:
            return ""
        if value[0] in ("=", "+", "-", "@", "\t", "\r"):
            return "'" + value
        stripped = value.lstrip()
        if stripped and stripped[0] in ("=", "+", "-", "@"):
            return "'" + value
        return value
    return value


def format_dt(dt):
    if not dt:
        return ""
    try:
        local = timezone.localtime(dt)
        return local.strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return str(dt)


def format_duration(seconds):
    sec = max(0, int(seconds or 0))
    m, s = divmod(sec, 60)
    h, m = divmod(m, 60)
    if h > 0:
        return f"{h:02d}:{m:02d}:{s:02d}"
    return f"{m:02d}:{s:02d}"


class StreamingBuffer:
    def write(self, value):
        return value


# -----------------------------------------------------------------------------
# Lead Directory Export Definitions
# -----------------------------------------------------------------------------

LEAD_EXPORT_COLUMNS = [
    ("Lead ID", "id"),
    ("Name", "name"),
    ("Phone", "phone"),
    ("Email", "email"),
    ("Status", "status"),
    ("Assigned Caller", "assigned_caller"),
    ("Caller Username", "caller_username"),
    ("Import Batch", "batch"),
    ("Source", "source"),
    ("Service", "service"),
    ("Preferred Course", "preferred_course"),
    ("Location", "location"),
    ("College", "college"),
    ("NEET Status", "neet_status"),
    ("PCB %", "pcb_percentage"),
    ("Preferred Intake", "preferred_intake"),
    ("Created At", "created_at"),
    ("Updated At", "updated_at"),
    ("Notes", "notes"),
]


def lead_to_row(lead):
    caller_name = ""
    caller_user = ""
    if lead.assigned_caller:
        caller_name = lead.assigned_caller.get_full_name() or lead.assigned_caller.username
        caller_user = lead.assigned_caller.username

    batch_display = "Manual / Unbatched"
    if lead.import_batch:
        batch_display = f"{lead.import_batch.filename} (#{lead.import_batch.pk})"

    service_name = ""
    if lead.service_type:
        service_name = lead.service_type.name
    elif lead.service:
        service_name = lead.service

    return [
        lead.pk,
        sanitize_cell(lead.name),
        sanitize_cell(lead.phone),
        sanitize_cell(lead.email),
        sanitize_cell(lead.get_status_display() or lead.status),
        sanitize_cell(caller_name or "Unassigned"),
        sanitize_cell(caller_user),
        sanitize_cell(batch_display),
        sanitize_cell(lead.source),
        sanitize_cell(service_name),
        sanitize_cell(lead.preferred_course_label),
        sanitize_cell(lead.location),
        sanitize_cell(lead.college),
        sanitize_cell(lead.neet_status),
        lead.pcb_percentage if lead.pcb_percentage is not None else "",
        sanitize_cell(lead.preferred_intake),
        format_dt(lead.created_at),
        format_dt(lead.updated_at),
        sanitize_cell(lead.notes),
    ]


def export_leads_csv(queryset):
    """Streams filtered leads as a CSV file with formula-injection protection and UTF-8 BOM."""
    timestamp = timezone.localtime().strftime("%Y%m%d_%H%M%S")
    filename = f"lead_directory_{timestamp}.csv"

    def row_generator():
        yield "\ufeff"  # UTF-8 BOM for Microsoft Excel compatibility
        buffer = io.StringIO()
        writer = csv.writer(buffer)

        # Header
        writer.writerow([col[0] for col in LEAD_EXPORT_COLUMNS])
        yield buffer.getvalue()
        buffer.seek(0)
        buffer.truncate(0)

        # Optimized iteration
        leads_qs = queryset.select_related("assigned_caller", "import_batch", "service_type")
        for lead in leads_qs.iterator(chunk_size=1000):
            writer.writerow(lead_to_row(lead))
            yield buffer.getvalue()
            buffer.seek(0)
            buffer.truncate(0)

    response = StreamingHttpResponse(row_generator(), content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def export_leads_excel(queryset):
    """Exports filtered leads as an Excel (.xlsx) file using openpyxl."""
    timestamp = timezone.localtime().strftime("%Y%m%d_%H%M%S")
    filename = f"lead_directory_{timestamp}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.views.sheetView[0].showGridLines = True

    # Styling
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    regular_font = Font(name="Segoe UI", size=10)
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    # Write headers
    headers = [col[0] for col in LEAD_EXPORT_COLUMNS]
    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        cell.border = thin_border
    ws.row_dimensions[1].height = 26

    # Write data
    leads_qs = queryset.select_related("assigned_caller", "import_batch", "service_type")
    row_idx = 2
    for lead in leads_qs.iterator(chunk_size=1000):
        row_data = lead_to_row(lead)
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 20
        for col_num in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_num)
            cell.font = regular_font
            cell.border = thin_border
            # Align IDs and numeric fields right, others left
            if col_num in (1, 15):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        row_idx += 1

    # Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    response = HttpResponse(
        bio.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


# -----------------------------------------------------------------------------
# Call Activity Export Definitions
# -----------------------------------------------------------------------------

CALL_EXPORT_COLUMNS = [
    ("Call ID", "id"),
    ("Started At", "started_at"),
    ("Ended At", "ended_at"),
    ("Duration (sec)", "duration_seconds"),
    ("Duration (mm:ss)", "duration_formatted"),
    ("Caller Name", "caller_name"),
    ("Caller Username", "caller_username"),
    ("Lead Name", "lead_name"),
    ("Lead Phone", "lead_phone"),
    ("Dialed Phone", "phone_number"),
    ("Outcome", "outcome"),
    ("Selected Course", "selected_course"),
    ("Expected Admission Year", "expected_admission_year"),
    ("Has Recording", "has_recording"),
    ("Has Follow-up", "has_followup"),
    ("Follow-up Scheduled At", "followup_scheduled_at"),
    ("Follow-up Status", "followup_status"),
    ("Lead Batch", "lead_batch"),
    ("Lead Source", "lead_source"),
    ("Notes", "notes"),
]


def call_to_row(call):
    caller_name = ""
    caller_user = ""
    if call.caller:
        caller_name = call.caller.get_full_name() or call.caller.username
        caller_user = call.caller.username

    lead_name = ""
    lead_phone = ""
    lead_batch = ""
    lead_source = ""
    if call.lead:
        lead_name = call.lead.name
        lead_phone = call.lead.phone
        lead_source = call.lead.source or ""
        if call.lead.import_batch:
            lead_batch = f"{call.lead.import_batch.filename} (#{call.lead.import_batch.pk})"
        else:
            lead_batch = "Manual / Unbatched"

    has_recording = "Yes" if hasattr(call, "recording") and call.recording else "No"

    has_followup = "No"
    followup_time = ""
    followup_status = ""
    if hasattr(call, "followup") and call.followup:
        has_followup = "Yes"
        followup_time = format_dt(call.followup.scheduled_at)
        followup_status = call.followup.get_status_display() or call.followup.status

    return [
        call.pk,
        format_dt(call.started_at),
        format_dt(call.ended_at),
        call.duration_seconds,
        format_duration(call.duration_seconds),
        sanitize_cell(caller_name),
        sanitize_cell(caller_user),
        sanitize_cell(lead_name or "External call"),
        sanitize_cell(lead_phone),
        sanitize_cell(call.phone_number),
        sanitize_cell(call.get_outcome_display() or call.outcome or "No outcome"),
        sanitize_cell(call.selected_course_label),
        call.expected_admission_year or "",
        has_recording,
        has_followup,
        followup_time,
        sanitize_cell(followup_status),
        sanitize_cell(lead_batch),
        sanitize_cell(lead_source),
        sanitize_cell(call.notes),
    ]


def export_calls_csv(queryset):
    """Streams filtered call activity as a CSV file with formula-injection protection and UTF-8 BOM."""
    timestamp = timezone.localtime().strftime("%Y%m%d_%H%M%S")
    filename = f"call_activity_{timestamp}.csv"

    def row_generator():
        yield "\ufeff"  # UTF-8 BOM
        buffer = io.StringIO()
        writer = csv.writer(buffer)

        writer.writerow([col[0] for col in CALL_EXPORT_COLUMNS])
        yield buffer.getvalue()
        buffer.seek(0)
        buffer.truncate(0)

        calls_qs = queryset.select_related(
            "lead", "lead__import_batch", "caller", "recording", "followup"
        )
        for call in calls_qs.iterator(chunk_size=1000):
            writer.writerow(call_to_row(call))
            yield buffer.getvalue()
            buffer.seek(0)
            buffer.truncate(0)

    response = StreamingHttpResponse(row_generator(), content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def export_calls_excel(queryset):
    """Exports filtered call activity as an Excel (.xlsx) file using openpyxl."""
    timestamp = timezone.localtime().strftime("%Y%m%d_%H%M%S")
    filename = f"call_activity_{timestamp}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Call Activity"
    ws.views.sheetView[0].showGridLines = True

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    regular_font = Font(name="Segoe UI", size=10)
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    headers = [col[0] for col in CALL_EXPORT_COLUMNS]
    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        cell.border = thin_border
    ws.row_dimensions[1].height = 26

    calls_qs = queryset.select_related(
        "lead", "lead__import_batch", "caller", "recording", "followup"
    )
    row_idx = 2
    for call in calls_qs.iterator(chunk_size=1000):
        row_data = call_to_row(call)
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 20
        for col_num in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_num)
            cell.font = regular_font
            cell.border = thin_border
            if col_num in (1, 4, 13):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        row_idx += 1

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    response = HttpResponse(
        bio.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response

