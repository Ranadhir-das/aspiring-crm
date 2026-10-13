import csv
import io
from datetime import timedelta
import openpyxl
from django.test import TestCase
from django.utils import timezone
from apps.accounts.models import User
from apps.calls.models import Call
from apps.followups.models import FollowUp
from apps.leads.models import Lead, LeadImportBatch
from apps.web.export_utils import sanitize_cell


class ReportingAndExportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.manager = User.objects.create_user('manager_rep', password='test-pass', role=User.Role.MANAGER)
        cls.caller_a = User.objects.create_user('caller_a', password='test-pass', role=User.Role.CALLER)
        cls.caller_b = User.objects.create_user('caller_b', password='test-pass', role=User.Role.CALLER)

        cls.batch1 = LeadImportBatch.objects.create(filename='batch_alpha.csv', imported_by=cls.manager)
        cls.batch2 = LeadImportBatch.objects.create(filename='batch_beta.csv', imported_by=cls.manager)

        # Lead 1: assigned to caller_a, batch1, INTERESTED
        cls.lead1 = Lead.objects.create(
            name='Alice Wonderland',
            phone='9876543210',
            email='alice@example.com',
            status=Lead.Status.INTERESTED,
            source='Website',
            service='BTECH',
            assigned_caller=cls.caller_a,
            import_batch=cls.batch1,
        )

        # Lead 2: assigned to caller_b, batch2, CALLED
        cls.lead2 = Lead.objects.create(
            name='Bob Builder',
            phone='9876543211',
            email='bob@example.com',
            status=Lead.Status.CALLED,
            source='Referral',
            service='MBA',
            assigned_caller=cls.caller_b,
            import_batch=cls.batch2,
        )

        # Lead 3: Unassigned, manual/unbatched, formula injection name
        cls.lead3 = Lead.objects.create(
            name='=cmd|/c calc!A0',
            phone='9876543212',
            email='danger@example.com',
            status=Lead.Status.NOT_INTERESTED,
            source='Walk-in',
            service='BTECH',
            assigned_caller=None,
            import_batch=None,
        )

        # Calls
        now = timezone.now()
        cls.call1 = Call.objects.create(
            lead=cls.lead1,
            caller=cls.caller_a,
            started_at=now - timedelta(days=2),
            duration_seconds=45,
            outcome=Call.Outcome.INTERESTED,
            notes='Positive conversation about BTECH',
        )
        cls.call2 = Call.objects.create(
            lead=cls.lead1,
            caller=cls.caller_a,
            started_at=now - timedelta(days=1),
            duration_seconds=150,
            outcome=Call.Outcome.INTERESTED,
            notes='Second call with Alice',
        )
        cls.call3 = Call.objects.create(
            lead=cls.lead2,
            caller=cls.caller_b,
            started_at=now - timedelta(hours=3),
            duration_seconds=20,
            outcome=Call.Outcome.BUSY,
            notes='Bob was busy',
        )
        cls.call_ext = Call.objects.create(
            caller=cls.caller_a,
            phone_number='9998887776',
            started_at=now - timedelta(minutes=30),
            duration_seconds=15,
            outcome=Call.Outcome.NO_ANSWER,
            notes='External inquiry',
        )

        # FollowUps
        cls.fu1 = FollowUp.objects.create(
            lead=cls.lead1,
            caller=cls.caller_a,
            scheduled_at=now + timedelta(days=1),
            status=FollowUp.Status.PENDING,
            notes='Follow-up on admission forms',
        )
        cls.fu2 = FollowUp.objects.create(
            lead=cls.lead2,
            caller=cls.caller_b,
            scheduled_at=now - timedelta(days=1),
            status=FollowUp.Status.PENDING,
            notes='Overdue follow-up',
        )

    def test_formula_injection_sanitization(self):
        self.assertEqual(sanitize_cell("=1+1"), "'=1+1")
        self.assertEqual(sanitize_cell("+cmd"), "'+cmd")
        self.assertEqual(sanitize_cell("-cmd"), "'-cmd")
        self.assertEqual(sanitize_cell("@SUM(A1:A5)"), "'@SUM(A1:A5)")
        self.assertEqual(sanitize_cell("\tTabbed"), "'\tTabbed")
        self.assertEqual(sanitize_cell("\rReturn"), "'\rReturn")
        self.assertEqual(sanitize_cell("Normal Text"), "Normal Text")
        self.assertEqual(sanitize_cell(12345), 12345)
        self.assertEqual(sanitize_cell(None), "")

    def test_lead_directory_filtering_search_and_status(self):
        self.client.force_login(self.manager)

        # Search Alice
        res = self.client.get('/leads/?q=Alice')
        self.assertEqual(res.status_code, 200)
        pks = [l.pk for l in res.context['records']]
        self.assertIn(self.lead1.pk, pks)
        self.assertNotIn(self.lead2.pk, pks)

        # Search Phone
        res = self.client.get('/leads/?q=9876543211')
        pks = [l.pk for l in res.context['records']]
        self.assertEqual(pks, [self.lead2.pk])

        # Filter Status
        res = self.client.get(f'/leads/?status={Lead.Status.INTERESTED}')
        pks = [l.pk for l in res.context['records']]
        self.assertEqual(pks, [self.lead1.pk])

    def test_lead_directory_batch_and_owner_filters(self):
        self.client.force_login(self.manager)

        # Batch 1
        res = self.client.get(f'/leads/?batch={self.batch1.pk}')
        pks = [l.pk for l in res.context['records']]
        self.assertEqual(pks, [self.lead1.pk])

        # Batch unbatched / none
        res = self.client.get('/leads/?batch=none')
        pks = [l.pk for l in res.context['records']]
        self.assertEqual(pks, [self.lead3.pk])

        # Owner unassigned
        res = self.client.get('/leads/?owner=unassigned')
        pks = [l.pk for l in res.context['records']]
        self.assertEqual(pks, [self.lead3.pk])

        # Owner caller_a
        res = self.client.get(f'/leads/?owner={self.caller_a.pk}')
        pks = [l.pk for l in res.context['records']]
        self.assertEqual(pks, [self.lead1.pk])

    def test_lead_directory_subquery_filters_no_duplicate_rows(self):
        self.client.force_login(self.manager)

        # Lead1 has 2 calls. Filtering by last call date should not yield duplicate Lead1 rows.
        today_str = timezone.now().date().strftime('%Y-%m-%d')
        res = self.client.get(f'/leads/?last_call_to={today_str}')
        pks = [l.pk for l in res.context['records']]
        # Lead1 and Lead2 have calls
        self.assertEqual(len(pks), len(set(pks)), "No duplicate lead rows returned")
        self.assertIn(self.lead1.pk, pks)
        self.assertIn(self.lead2.pk, pks)
        self.assertNotIn(self.lead3.pk, pks)

        # Follow-up filter: upcoming
        res = self.client.get('/leads/?followup_status=upcoming')
        pks = [l.pk for l in res.context['records']]
        self.assertEqual(pks, [self.lead1.pk])

        # Follow-up filter: overdue
        res = self.client.get('/leads/?followup_status=overdue')
        pks = [l.pk for l in res.context['records']]
        self.assertEqual(pks, [self.lead2.pk])

        # Follow-up filter: no_followup
        res = self.client.get('/leads/?followup_status=no_followup')
        pks = [l.pk for l in res.context['records']]
        self.assertIn(self.lead3.pk, pks)
        self.assertNotIn(self.lead1.pk, pks)
        self.assertNotIn(self.lead2.pk, pks)

    def test_lead_directory_date_validation_error(self):
        self.client.force_login(self.manager)

        # Invalid format
        res = self.client.get('/leads/?start_date=invalid-date')
        self.assertEqual(res.status_code, 200)
        self.assertIsNotNone(res.context['date_error'])
        self.assertEqual(len(res.context['records']), 0)

        # start_date > end_date
        res = self.client.get('/leads/?start_date=2026-10-15&end_date=2026-10-01')
        self.assertEqual(res.status_code, 200)
        self.assertIn("Start date cannot be after end date", res.context['date_error'])
        self.assertEqual(len(res.context['records']), 0)

    def test_call_activity_filtering(self):
        self.client.force_login(self.manager)

        # Search by lead name
        res = self.client.get('/calls/?q=Alice')
        pks = [c.pk for c in res.context['records']]
        self.assertIn(self.call1.pk, pks)
        self.assertIn(self.call2.pk, pks)
        self.assertNotIn(self.call3.pk, pks)

        # Search by external phone
        res = self.client.get('/calls/?q=9998887776')
        pks = [c.pk for c in res.context['records']]
        self.assertEqual(pks, [self.call_ext.pk])

        # Filter by outcome
        res = self.client.get(f'/calls/?outcome={Call.Outcome.BUSY}')
        pks = [c.pk for c in res.context['records']]
        self.assertEqual(pks, [self.call3.pk])

        # Duration preset: under 30s
        res = self.client.get('/calls/?duration_preset=under_30')
        pks = [c.pk for c in res.context['records']]
        self.assertIn(self.call3.pk, pks)
        self.assertIn(self.call_ext.pk, pks)
        self.assertNotIn(self.call2.pk, pks)

        # Duration preset: over 2m (>120s)
        res = self.client.get('/calls/?duration_preset=over_120')
        pks = [c.pk for c in res.context['records']]
        self.assertEqual(pks, [self.call2.pk])

        # Duration numeric range
        res = self.client.get('/calls/?duration_min=40&duration_max=50')
        pks = [c.pk for c in res.context['records']]
        self.assertEqual(pks, [self.call1.pk])

        # Batch filter via lead
        res = self.client.get(f'/calls/?batch={self.batch2.pk}')
        pks = [c.pk for c in res.context['records']]
        self.assertEqual(pks, [self.call3.pk])

    def _get_response_bytes(self, response):
        if hasattr(response, 'streaming_content'):
            return b"".join(response.streaming_content)
        return response.content

    def test_lead_csv_export_and_sanitization(self):
        self.client.force_login(self.manager)

        res = self.client.get('/leads/?export=csv')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res['Content-Type'], 'text/csv; charset=utf-8')
        self.assertIn('attachment; filename="lead_directory_', res['Content-Disposition'])

        content = self._get_response_bytes(res).decode('utf-8-sig')  # strip BOM if present
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)

        headers = rows[0]
        self.assertIn('Lead ID', headers)
        self.assertIn('Name', headers)
        self.assertIn('Phone', headers)
        self.assertIn('Status', headers)

        data_rows = rows[1:]
        self.assertEqual(len(data_rows), 3)

        # Check sanitization of formula injection in lead3
        lead3_row = [r for r in data_rows if r[2] == '9876543212'][0]
        self.assertTrue(lead3_row[1].startswith("'=cmd"), f"Formula escaped with ': {lead3_row[1]}")

        # Check export preserves active filter
        res_filtered = self.client.get('/leads/?q=Alice&export=csv')
        filtered_rows = list(csv.reader(io.StringIO(self._get_response_bytes(res_filtered).decode('utf-8-sig'))))[1:]
        self.assertEqual(len(filtered_rows), 1)
        self.assertEqual(filtered_rows[0][1], 'Alice Wonderland')

    def test_lead_excel_export(self):
        self.client.force_login(self.manager)

        res = self.client.get('/leads/?export=xlsx')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('attachment; filename="lead_directory_', res['Content-Disposition'])

        wb = openpyxl.load_workbook(io.BytesIO(self._get_response_bytes(res)))
        self.assertIn('Leads', wb.sheetnames)
        sheet = wb['Leads']

        headers = [cell.value for cell in sheet[1]]
        self.assertIn('Lead ID', headers)
        self.assertIn('Name', headers)
        self.assertIn('Assigned Caller', headers)

        # Check rows
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 3)

        # Formula sanitization check
        names = [r[1] for r in rows]
        self.assertTrue(any(str(n).startswith("'=") for n in names))

    def test_call_csv_and_excel_export(self):
        self.client.force_login(self.manager)

        # CSV
        res_csv = self.client.get('/calls/?export=csv')
        self.assertEqual(res_csv.status_code, 200)
        self.assertEqual(res_csv['Content-Type'], 'text/csv; charset=utf-8')
        self.assertIn('attachment; filename="call_activity_', res_csv['Content-Disposition'])
        csv_rows = list(csv.reader(io.StringIO(self._get_response_bytes(res_csv).decode('utf-8-sig'))))
        self.assertIn('Duration (sec)', csv_rows[0])
        self.assertEqual(len(csv_rows) - 1, 4)

        # Excel
        res_xlsx = self.client.get('/calls/?export=xlsx')
        self.assertEqual(res_xlsx.status_code, 200)
        self.assertIn('attachment; filename="call_activity_', res_xlsx['Content-Disposition'])
        wb = openpyxl.load_workbook(io.BytesIO(self._get_response_bytes(res_xlsx)))
        self.assertIn('Call Activity', wb.sheetnames)
        sheet = wb['Call Activity']
        self.assertEqual(sheet.max_row, 5)  # 1 header + 4 rows

    def test_caller_data_isolation_on_exports_and_views(self):
        # Caller A should ONLY see and export lead1 and calls (call1, call2, call_ext)
        self.client.force_login(self.caller_a)

        # HTML leads
        res_leads = self.client.get('/leads/')
        self.assertEqual(res_leads.status_code, 200)
        lead_pks = [l.pk for l in res_leads.context['records']]
        self.assertEqual(lead_pks, [self.lead1.pk])

        # CSV leads export
        res_csv = self.client.get('/leads/?export=csv')
        self.assertEqual(res_csv.status_code, 200)
        csv_rows = list(csv.reader(io.StringIO(self._get_response_bytes(res_csv).decode('utf-8-sig'))))[1:]
        self.assertEqual(len(csv_rows), 1)
        self.assertEqual(csv_rows[0][0], str(self.lead1.pk))

        # Excel leads export
        res_xlsx = self.client.get('/leads/?export=xlsx')
        wb = openpyxl.load_workbook(io.BytesIO(self._get_response_bytes(res_xlsx)))
        rows = list(wb['Leads'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][0], self.lead1.pk)

        # HTML calls
        res_calls = self.client.get('/calls/')
        self.assertEqual(res_calls.status_code, 200)
        call_pks = [c.pk for c in res_calls.context['records']]
        self.assertNotIn(self.call3.pk, call_pks)  # caller_b call is NOT visible

        # CSV calls export
        res_calls_csv = self.client.get('/calls/?export=csv')
        call_csv_rows = list(csv.reader(io.StringIO(self._get_response_bytes(res_calls_csv).decode('utf-8-sig'))))[1:]
        exported_call_ids = [int(r[0]) for r in call_csv_rows]
        self.assertNotIn(self.call3.pk, exported_call_ids)
        self.assertIn(self.call1.pk, exported_call_ids)

        # Hostile filter parameter injection: Caller A attempts to filter by Caller B's ID or unassigned
        res_hostile_leads = self.client.get(f'/leads/?owner={self.caller_b.pk}')
        self.assertEqual([l.pk for l in res_hostile_leads.context['records']], [self.lead1.pk])
        res_hostile_leads_csv = self.client.get(f'/leads/?owner={self.caller_b.pk}&export=csv')
        hostile_lead_ids = [int(r[0]) for r in list(csv.reader(io.StringIO(self._get_response_bytes(res_hostile_leads_csv).decode('utf-8-sig'))))[1:]]
        self.assertEqual(hostile_lead_ids, [self.lead1.pk])

        res_hostile_calls = self.client.get(f'/calls/?caller={self.caller_b.pk}')
        self.assertNotIn(self.call3.pk, [c.pk for c in res_hostile_calls.context['records']])
        res_hostile_calls_csv = self.client.get(f'/calls/?caller={self.caller_b.pk}&export=csv')
        hostile_call_ids = [int(r[0]) for r in list(csv.reader(io.StringIO(self._get_response_bytes(res_hostile_calls_csv).decode('utf-8-sig'))))[1:]]
        self.assertNotIn(self.call3.pk, hostile_call_ids)

    def test_export_includes_all_records_regardless_of_pagination(self):
        self.client.force_login(self.manager)

        # Create 25 additional leads so total exceeds the page size of 20
        extra_leads = [
            Lead(name=f"Bulk Lead {i}", phone=f"90000000{i:02d}", status=Lead.Status.CALLED)
            for i in range(25)
        ]
        Lead.objects.bulk_create(extra_leads)
        total_leads = Lead.objects.count()
        self.assertGreater(total_leads, 20)

        # HTML page only displays page 1 (20 records)
        res_html = self.client.get('/leads/')
        self.assertEqual(res_html.status_code, 200)
        self.assertEqual(len(res_html.context['records']), 20)
        self.assertEqual(res_html.context['records'].paginator.count, total_leads)

        # CSV export downloads ALL records (not just page 1)
        res_csv = self.client.get('/leads/?export=csv')
        self.assertEqual(res_csv.status_code, 200)
        csv_rows = list(csv.reader(io.StringIO(self._get_response_bytes(res_csv).decode('utf-8-sig'))))[1:]
        self.assertEqual(len(csv_rows), total_leads)

        # Excel export includes ALL records
        res_xlsx = self.client.get('/leads/?export=xlsx')
        self.assertEqual(res_xlsx.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(self._get_response_bytes(res_xlsx)))
        excel_rows = list(wb['Leads'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(excel_rows), total_leads)
